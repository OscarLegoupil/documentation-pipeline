"""Reproducible reference document and narrowly scoped Word finishing."""
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

from .common import Failure


def build_reference(path, branding=None):
    from docx import Document
    from docx.shared import Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document()
    colors = branding or {"navy": "17324D", "teal": "007F82"}
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(20)
    section.header_distance = section.footer_distance = Mm(9)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name in ("Body Text", "First Paragraph", "Compact", "Source Note", "Contents Entry", "Caption", "Image Caption", "Table Caption"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[name]
        style.base_style = normal
        style.font.size = Pt(9 if name in ("Source Note", "Caption", "Image Caption", "Table Caption") else 11)
    for level in range(1, 5):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt([20, 15, 12, 11][level - 1])
        style.font.color.rgb = RGBColor.from_string(colors["navy"] if level < 3 else colors["teal"])
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(15)
        style.paragraph_format.space_after = Pt(7)
    for name in ("Source Code", "Verbatim Char"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH if name == "Source Code" else WD_STYLE_TYPE.CHARACTER)
        doc.styles[name].font.name = "Consolas"
        doc.styles[name].font.size = Pt(9)
    doc.styles["Title"].font.size = Pt(32)
    doc.styles["Title"].font.color.rgb = RGBColor.from_string(colors["navy"])
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Documentation reference style"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(path)


def finish(path):
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(path)
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            p.paragraph_format.page_break_before = True
        if p.style.name in ("Caption", "Image Caption", "Table Caption"):
            p.paragraph_format.keep_with_next = False
        p.paragraph_format.widow_control = True
        for run in p.runs:
            # Break opportunities in long paths; actual source/evidence paths remain exact in JSON.
            if run.style and run.style.name in ("Verbatim Char", "Code"):
                if len(run.text) > 35:
                    run.text = run.text.replace("/", "/\u200b").replace("\\", "\\\u200b").replace("_", "_\u200b")
    for table in doc.tables:
        table.autofit = False
        columns = len(table.columns)
        width = Mm(170 / max(1, columns))
        for col in table.columns:
            col.width = width
        for idx, row in enumerate(table.rows):
            for cell in row.cells:
                cell.width = width
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    for run in p.runs:
                        run.font.size = Pt(9.5)
                        if len(run.text) > 35:
                            run.text = run.text.replace("/", "/\u200b").replace("_", "_\u200b")
            if idx == 0:
                trpr = row._tr.get_or_add_trPr()
                if trpr.find(qn("w:tblHeader")) is None:
                    trpr.append(OxmlElement("w:tblHeader"))
                for cell in row.cells:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E8F0F4")
                    cell._tc.get_or_add_tcPr().append(shade)
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
                        for r in p.runs:
                            r.bold = True
    for picture in doc.inline_shapes:
        if picture.width > Mm(170):
            ratio = Mm(170) / picture.width
            picture.width = Mm(170)
            picture.height = int(picture.height * ratio)
        if picture.height > Mm(215):
            ratio = Mm(215) / picture.height
            picture.height = Mm(215)
            picture.width = int(picture.width * ratio)
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.save(path)
    return inspect_docx(path)


def inspect_docx(path):
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
        q = lambda name: "{" + ns["w"] + "}" + name
        bookmarks = {n.attrib[q("name")] for n in root.findall(".//w:bookmarkStart", ns)}
        anchors = [n.attrib[q("anchor")] for n in root.findall(".//w:hyperlink", ns) if q("anchor") in n.attrib]
        missing = sorted(set(anchors) - bookmarks)
        if missing:
            raise Failure(f"DOCX hyperlinks target missing bookmarks: {missing}")
        if not anchors:
            raise Failure("DOCX has no visible linked contents/evidence hyperlinks")
        texts = [n.text or "" for n in root.findall(".//w:t", ns)]
        if "Contents" not in texts:
            raise Failure("DOCX visible contents heading missing")
        return {"bookmarks": len(bookmarks), "internal_hyperlinks": len(anchors), "tables": len(root.findall(".//w:tbl", ns)), "repeating_headers": len(root.findall(".//w:tblHeader", ns)), "embedded_images": len([n for n in z.namelist() if n.startswith("word/media/")]), "drawings": len(root.findall(".//w:drawing", ns)), "headings": len([n for n in root.findall(".//w:pStyle", ns) if n.attrib.get(q("val"), "").startswith("Heading")]), "missing_bookmarks": missing}
