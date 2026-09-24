import copy
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from support import Engine, fixture, finish_fixture, accept_next, artifact, config, tool_path, SKILL
from pipeline.common import Failure, atomic, digest, read_json
from pipeline.contracts import registry, validate_result
from pipeline.render import export, check_markdown, diagram_assets
from pipeline import render
from pipeline.style import inspect_docx


class IntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        for name in ("pandoc", "dot"):
            if not tool_path(name):
                raise RuntimeError(f"Integration tests REQUIRE {name}; install explicitly. Run test_core.py for pure Python tests.")

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="docgen export ü ")
        self.root = fixture(root=self.temp.name)
        self.engine = Engine(self.root)
        self.session = self.engine.prepare()["session"]

    def tearDown(self):
        self.temp.cleanup()

    def complete(self, layout=False):
        status = finish_fixture(self.engine, self.session, layout)
        self.assertTrue(status["pipeline_complete"])
        self.assertTrue(self.engine.validate()["valid"])
        return status

    def test_export_docx_structure_and_repeat_without_model_or_network(self):
        with patch("socket.socket", side_effect=AssertionError("Network forbidden")):
            self.complete(layout=True)
            before = {i: t.get("accepted_hash") for i, t in self.engine.state["tasks"].items()}
            build = export(self.engine)
            self.assertTrue(Path(build["docx"]).is_file())
            self.assertGreater(build["docx_checks"]["embedded_images"], 0)
            self.assertGreater(build["docx_checks"]["tables"], 3)
            self.assertEqual(build["docx_checks"]["tables"], build["docx_checks"]["repeating_headers"])
            self.assertFalse(build["docx_checks"]["missing_bookmarks"])
            heading_map = read_json(Path(build["directory"]) / "heading-map.json")
            self.assertEqual(len(heading_map), len({x["bookmark"] for x in heading_map}))
            self.assertFalse(any(".0." in x["text"] for x in heading_map))
            self.assertEqual(build["visual_qa"], "pending")
            md = Path(build["markdown"]).read_text("utf-8")
            self.assertIn("Київ", md)
            self.assertIn("Figure 1", md)
            with zipfile.ZipFile(build["docx"]) as z:
                xml = z.read("word/document.xml").decode()
                self.assertIn("w:pageBreakBefore", xml)
                self.assertIn("w:bookmarkStart", xml)
                self.assertIn("w:anchor", xml)
                self.assertNotIn("TOC ", xml)
            self.assertTrue(self.engine.validate()["valid"], "Generated outputs must self-exclude")
            second = export(self.engine)
            self.assertEqual(Path(second["markdown"]).read_text("utf-8"), md)
            self.assertEqual(before, {i: t.get("accepted_hash") for i, t in self.engine.state["tasks"].items()})

    def test_strict_partial_and_broken_tools(self):
        with self.assertRaises(Failure) as caught:
            export(self.engine)
        self.assertEqual(caught.exception.code, 6)
        draft = export(self.engine, partial=True)
        self.assertEqual(Path(draft["docx"]).name, "DRAFT-INCOMPLETE.docx")
        self.assertIn("DRAFT", Path(draft["markdown"]).read_text("utf-8"))
        self.complete()
        old = self.engine.config["executables"]["dot"]
        self.engine.config["executables"]["dot"] = str(self.root / "absent-renderer.exe")
        with self.assertRaises(Failure):
            export(self.engine)
        self.engine.config["executables"]["dot"] = old
        self.engine.config["executables"]["pandoc"] = str(self.root / "absent-pandoc.exe")
        with self.assertRaises(Failure):
            export(self.engine)

    def test_incremental_chapters_and_preservation_on_export_failure(self):
        from docx import Document
        root = fixture("web-service", root=self.root / "service fixture")
        engine = Engine(root)
        session = engine.prepare()["session"]
        progress_path = root / engine.config["output_dir"] / engine.sid / "documentation-progress.docx"
        blocked_author = None
        storage_review = None
        # Hand-authored artifacts: block API, complete ledger (with diagram), leave storage pending.
        with patch("socket.socket", side_effect=AssertionError("Network forbidden")):
            for _ in range(30):
                receipt = engine.next(session)
                task = engine.state["tasks"][receipt["task_id"]]
                result = artifact(engine, task)
                if task["kind"] == "review":
                    author = engine.state["tasks"][task["deps"][0]]
                    subsystem = author["payload"]["subsystem"]
                    if subsystem == "storage":
                        storage_review = task
                        break
                    if subsystem == "api":
                        blocked_author = author["id"]
                        result["findings"] = [{"id": task["id"] + ".blocker", "severity": "blocker", "blocks": [task["review_blocks"][0]], "facts": [], "evidence": [], "message": "Hand-authored fixture blocker."}]
                atomic(Path(receipt["result"]), result)
                engine.accept(task["id"], session)
                if task["kind"] != "review":
                    self.assertFalse(progress_path.exists(), "Accepted writing alone must not trigger a preview")
            self.assertIsNotNone(storage_review)
            self.assertIsNotNone(blocked_author)
            self.assertTrue(progress_path.is_file(), engine.state.get("progress_error"))
            preview = copy.deepcopy(engine.state["progress_preview"])
            self.assertEqual(preview["subsystems"], {"completed": 1, "pending": 1, "blocked": 1})
            progress = read_json(Path(preview["directory"]) / "progress.json")
            self.assertEqual(progress["completed"], ["ledger"])
            doc = Document(progress_path)
            text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("Incomplete draft", text)
            self.assertIn("Subsystem progress", text)
            self.assertGreater(len(doc.inline_shapes), 0)
            headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
            self.assertTrue(any("Ledger" in h for h in headings))
            self.assertFalse(any("Api" in h or "Storage" in h for h in headings))
            cells = " ".join(c.text for table in doc.tables for row in table.rows for c in row.cells)
            for expected in ("Completed", "Pending", "Blocked", "api", "storage"):
                self.assertIn(expected, cells)
            with self.assertRaises(Failure):
                export(engine)
            self.assertIsNone(engine.state.get("build"))
            previous_bytes = progress_path.read_bytes()
            before = {i: t.get("accepted_hash") for i, t in engine.state["tasks"].items() if t["status"] == "accepted"}
            # Fail late, after assembly/diagram work but before a completed Word document exists.
            real_run = render.run
            def failed_export(argv, *args, **kwargs):
                if "--to" in argv and argv[argv.index("--to") + 1] == "docx":
                    raise Failure("Fixture Pandoc DOCX failure", 3)
                return real_run(argv, *args, **kwargs)
            atomic(engine.task_dir(storage_review["id"]) / "result.json", artifact(engine, storage_review))
            with patch.object(render, "run", side_effect=failed_export):
                accepted = engine.accept(storage_review["id"], session)
            self.assertEqual(accepted["accepted"], storage_review["id"])
            self.assertIn("Fixture Pandoc DOCX failure", accepted["progress_error"])
            self.assertEqual(progress_path.read_bytes(), previous_bytes)
            self.assertEqual(engine.state["progress_preview"]["directory"], preview["directory"])
            self.assertEqual(engine.state["tasks"][storage_review["id"]]["status"], "accepted")
            # A locked destination also leaves the last successful bytes untouched.
            real_atomic = render.atomic
            def locked_destination(path, data):
                if Path(path) == progress_path:
                    raise PermissionError("Fixture Word document is open")
                return real_atomic(path, data)
            with patch.object(render, "atomic", side_effect=locked_destination):
                engine.next(session)
            self.assertEqual(progress_path.read_bytes(), previous_bytes)
            self.assertIn("Fixture Word document is open", engine.state["progress_error"])
            # Retry from disk: rebuild only, with all accepted tasks and analysis preserved.
            engine = Engine(root)
            engine.next(session)
            self.assertIsNone(engine.state["progress_error"])
            self.assertEqual(engine.state["progress_preview"]["subsystems"], {"completed": 2, "pending": 0, "blocked": 1})
            self.assertNotEqual(progress_path.read_bytes(), previous_bytes)
            for ident, accepted_hash in before.items():
                self.assertEqual(engine.state["tasks"][ident]["accepted_hash"], accepted_hash)
            current = progress_path.read_bytes()
            with patch.object(render, "render_document", side_effect=AssertionError("Unchanged preview must be reused")):
                engine.next(session)
            self.assertIsNone(engine.state["progress_error"])
            self.assertEqual(progress_path.read_bytes(), current)
            with self.assertRaises(Failure):
                export(engine)
            engine.repair(blocked_author, session)
            receipt = engine.next(session)
            repaired = artifact(engine, engine.state["tasks"][receipt["task_id"]])
            repaired["summary"] += " Fixture repair addressed the blocking finding."
            atomic(Path(receipt["result"]), repaired)
            engine.accept(receipt["task_id"], session)
            finish_fixture(engine, session)
            final = export(engine)
            self.assertEqual(Path(final["docx"]).name, "documentation.docx")
            self.assertNotEqual(Path(final["docx"]), progress_path)
            self.assertTrue(engine.validate()["valid"])
    def test_malformed_markdown_links_and_diagrams(self):
        self.complete()
        reg = registry(self.engine.results())
        fid = next(iter(reg["facts"]))
        for text in (f"Uncited paragraph.\n\nAnother [source](fact:{fid}).", f"A [source](fact:{fid}) ![remote](https://example.invalid/image.png)", f"A [source](fact:{fid}) and {{{{PLACEHOLDER}}}}", f"| Claim | Evidence |\n|---|---|\n| unsupported | missing |\n\n[source](fact:{fid})"):
            with self.assertRaises(Failure):
                check_markdown({"id": "test.block", "markdown": text, "facts": [fid]}, self.root, self.engine.config, reg)
        writer = next(t for t in self.engine.state["tasks"].values() if t["kind"] == "write" and t["status"] == "accepted")
        bad = copy.deepcopy(self.engine.result(writer))
        bad["diagrams"][0]["edges"][0]["target"] = "missing"
        others = [r for r in self.engine.results() if r["task_id"] != writer["id"]]
        with self.assertRaises(Failure):
            validate_result(self.root, self.engine.config, writer, bad, others, self.engine.state["manifest"], self.engine.allowed(writer))

    def test_semantic_blocker_repair_is_bounded_and_rereviewed(self):
        for _ in range(30):
            receipt = self.engine.next(self.session)
            task = self.engine.state["tasks"][receipt["task_id"]]
            result = artifact(self.engine, task)
            if task["kind"] == "review":
                result["findings"] = [{"id": task["id"] + ".finding1", "severity": "blocker", "blocks": [task["review_blocks"][0]], "facts": [task["review_facts"][0]], "evidence": [], "message": "Fixture review blocker: demonstrate bounded repair state."}]
            atomic(Path(receipt["result"]), result)
            self.engine.accept(task["id"], self.session)
            if task["kind"] == "review":
                author = task["deps"][0]
                break
        self.assertFalse(self.engine.validate()["valid"])
        self.assertTrue(self.engine.blockers())
        self.engine.repair(author, self.session)
        self.assertEqual(self.engine.state["tasks"][author]["cycles"], 1)
        receipt = self.engine.next(self.session)
        task = self.engine.state["tasks"][receipt["task_id"]]
        changed = artifact(self.engine, task)
        changed["summary"] += " Fixture repair addressed the finding."
        atomic(Path(receipt["result"]), changed)
        self.engine.accept(task["id"], self.session)
        self.complete()
        task = self.engine.state["tasks"][author]
        task["cycles"] = 2
        self.engine.save()
        with self.assertRaises(Failure) as caught:
            self.engine.repair(author, self.session)
        self.assertEqual(caught.exception.code, 6)

    def test_add_delete_rename_and_shared_contract_invalidation(self):
        self.complete()
        before = set(self.engine.state["active"])
        hashes = {i: self.engine.state["tasks"][i].get("accepted_hash") for i in before}
        self.engine.prepare(self.session)
        self.assertEqual(before, set(self.engine.state["active"]))
        self.assertEqual(hashes, {i: self.engine.state["tasks"][i].get("accepted_hash") for i in before})
        # Update a shared contract; every reconciled/chapter/review dependency is invalidated.
        path = self.root / "ledger.py"
        path.write_text(path.read_text("utf-8") + "\nCONTRACT_VERSION = 2\n", encoding="utf-8")
        self.assertFalse(self.engine.validate()["valid"])
        with self.assertRaises(Failure):
            export(self.engine, partial=True)
        self.engine.prepare(self.session)
        self.assertFalse(self.engine.status()["pipeline_complete"])
        old_downstream = {i for i in before if self.engine.state["tasks"][i]["kind"] in ("reconcile", "write", "review", "synthesize")}
        self.assertFalse(old_downstream & set(self.engine.state["active"]))
        self.complete()
        (self.root / "README.md").rename(self.root / "renamed.md")
        (self.root / "pyproject.toml").unlink()
        (self.root / "added.f90").write_text("program fixture\nend\n")
        self.engine.prepare(self.session)
        changes = self.engine.state["last_rescan"]["changed_paths"]
        self.assertTrue({"README.md", "renamed.md", "pyproject.toml", "added.f90"} <= set(changes))
        self.complete()

    def test_presentation_change_reuses_analysis(self):
        self.complete()
        before = {i: self.engine.state["tasks"][i]["accepted_hash"] for i in self.engine.state["active"]}
        config(self.root, title="Changed presentation only")
        self.engine = Engine(self.root)
        self.engine.prepare(self.session)
        self.assertEqual(before, {i: self.engine.state["tasks"][i]["accepted_hash"] for i in self.engine.state["active"]})
        build = export(self.engine)
        self.assertIn("Changed presentation only", Path(build["markdown"]).read_text("utf-8"))

    def test_service_and_generic_fallback_full_pipelines(self):
        for name in ("web-service", "mixed"):
            with self.subTest(fixture=name), tempfile.TemporaryDirectory(prefix="docgen nested ü ") as temp:
                root = fixture(name, temp)
                engine = Engine(root)
                session = engine.prepare()["session"]
                status = finish_fixture(engine, session)
                self.assertTrue(status["pipeline_complete"])
                self.assertTrue(engine.validate()["valid"])
                if name == "web-service":
                    reg = registry(engine.results())
                    self.assertTrue(any(r["source"] == "service/routes.ts" and r["target"] == "domain/orders.ts" and r["type"] == "call" for r in reg["relations"].values()))
                    self.assertTrue(any(f["basis"] == "declaration" for f in reg["facts"].values()))
                    self.assertTrue(any(f["status"] == "unknown" for f in reg["facts"].values()))
                    self.assertEqual(sum(len(r.get("entry_points", [])) for r in engine.results()), 3)
                    self.assertTrue(any(r.get("conflicts") for r in engine.results()))
                build = export(engine)
                self.assertTrue(Path(build["docx"]).exists())

    def test_installed_runtime_cli_paths_unicode(self):
        import importlib.util
        spec = importlib.util.spec_from_file_location("installer_under_test", SKILL / "install.py")
        installer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(installer)
        installer.install(SKILL, self.root)
        entry = self.root / ".github/skills/generate-documentation/docgen.py"
        import sys
        result = subprocess.run([sys.executable, str(entry), "--root", str(self.root), "doctor"], capture_output=True, text=True, encoding="utf-8", check=True, timeout=30)
        self.assertTrue(json.loads(result.stdout)["ok"])

    def test_accepted_artifact_corruption_is_detected_and_requeued(self):
        first = accept_next(self.engine, self.session)
        path = self.engine.task_dir(first) / "accepted.json"
        atomic(path, {"version": 1, "kind": "discover", "fake": True})
        with self.assertRaises(Failure):
            self.engine.status()
        self.engine.prepare(self.session)
        self.assertEqual(self.engine.state["tasks"][first]["status"], "pending")
        self.assertEqual(self.engine.next(self.session)["task_id"], first)

    def test_generated_large_fixture_resume_and_unchanged_packet_reuse(self):
        with tempfile.TemporaryDirectory(prefix="docgen large ") as tmp:
            root = Path(tmp)
            for i in range(60):
                (root / f"unit-{i:03d}.txt").write_text(f"unit {i:03d}\n", encoding="utf-8")
            (root / "oversized.txt").write_text("".join(f"source line {i:04d} contains explicit fixture text\n" for i in range(1000)), encoding="utf-8")
            (root / "empty.txt").write_text("")
            config(root)
            engine = Engine(root)
            session = engine.prepare()["session"]
            self.assertGreater(len(engine.state["packets"]), 5)
            completed = [accept_next(engine, session) for _ in range(4)]
            accepted_hashes = {i: engine.state["tasks"][i]["accepted_hash"] for i in completed}
            engine = Engine(root)
            session = engine.prepare(recover=True)["session"]
            self.assertEqual(accepted_hashes, {i: engine.state["tasks"][i]["accepted_hash"] for i in completed})
            for _ in range(250):
                ident = accept_next(engine, session)
                if engine.state["tasks"][ident]["kind"] == "review":
                    break
            writers = [engine.state["tasks"][i] for i in engine.state["active"] if engine.state["tasks"][i]["kind"] == "write"]
            self.assertGreater(len(writers), 1, "Large subsystem must exercise multiple chapter parts")
            self.assertIsNone(engine.state.get("progress_preview"), "One reviewed section is insufficient for a multi-section chapter")
            status = finish_fixture(engine, session)
            self.assertTrue(status["pipeline_complete"])
            self.assertEqual(engine.state["progress_preview"]["subsystems"]["completed"], 1)
            self.assertEqual(status["accepted_analysis_ranges"]["covered_lines"], status["eligible_lines"])
            self.assertEqual(status["accepted_analysis_ranges"]["covered_empty_files"], 1)
            old_analyses = {i for i in engine.state["active"] if engine.state["tasks"][i]["kind"] == "analyze"}
            (root / "unit-059.txt").write_text("unit 999\n", encoding="utf-8")
            engine.prepare(session)
            self.assertFalse(engine.state["progress_preview"]["current"], "An invalidated draft remains historical, not current")
            retained = old_analyses & set(engine.state["active"])
            self.assertTrue(retained)
            self.assertLess(len(retained), len(old_analyses))
            self.assertTrue(all(engine.state["tasks"][i]["status"] == "accepted" for i in retained))


if __name__ == "__main__":
    unittest.main()
